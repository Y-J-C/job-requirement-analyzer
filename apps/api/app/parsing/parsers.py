from io import BytesIO

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.parsing.contracts import DocumentParser, ImageOcr
from app.parsing.errors import DocumentError
from app.parsing.pdf_ocr import extract_pdf_text_with_ocr
from app.parsing.validation import DOCX_MEDIA_TYPE, MARKDOWN_MEDIA_TYPE, PDF_MEDIA_TYPE


def _normalize_text(text: str, *, max_chars: int) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "").strip()
    if len(normalized) > max_chars:
        raise DocumentError("extracted_text_too_large", "Extracted text exceeds the limit")
    if not normalized:
        raise DocumentError("no_extractable_text", "No extractable text was found")
    return normalized


class MarkdownParser:
    media_type = MARKDOWN_MEDIA_TYPE

    def parse(self, content: bytes, *, pdf_max_pages: int) -> str:
        del pdf_max_pages
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError as error:
            raise DocumentError("invalid_file_structure", "Markdown must be UTF-8") from error


class PdfParser:
    media_type = PDF_MEDIA_TYPE

    def parse(self, content: bytes, *, pdf_max_pages: int) -> str:
        try:
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise DocumentError("encrypted_document", "Encrypted PDF is not supported")
            if len(reader.pages) > pdf_max_pages:
                raise DocumentError("pdf_page_limit_exceeded", "PDF exceeds the page limit")
            return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
        except DocumentError:
            raise
        except (PdfReadError, OSError, ValueError) as error:
            raise DocumentError("document_parse_failed", "PDF could not be parsed") from error


class DocxParser:
    media_type = DOCX_MEDIA_TYPE

    def parse(self, content: bytes, *, pdf_max_pages: int) -> str:
        del pdf_max_pages
        try:
            document = Document(BytesIO(content))
            lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        lines.append("\t".join(cells))
            return "\n".join(lines)
        except (PackageNotFoundError, KeyError, OSError, ValueError) as error:
            raise DocumentError("document_parse_failed", "DOCX could not be parsed") from error


PARSERS: dict[str, DocumentParser] = {
    parser.media_type: parser for parser in (MarkdownParser(), PdfParser(), DocxParser())
}


def parse_document(
    content: bytes,
    media_type: str,
    *,
    max_chars: int,
    pdf_max_pages: int,
    image_ocr: ImageOcr | None = None,
    pdf_render_scale: float = 2,
    pdf_ocr_max_page_pixels: int = 20_000_000,
    pdf_ocr_max_total_pixels: int = 120_000_000,
) -> str:
    if media_type == PDF_MEDIA_TYPE and image_ocr is not None:
        return _normalize_text(
            extract_pdf_text_with_ocr(
                content,
                image_ocr=image_ocr,
                pdf_max_pages=pdf_max_pages,
                render_scale=pdf_render_scale,
                max_page_pixels=pdf_ocr_max_page_pixels,
                max_total_pixels=pdf_ocr_max_total_pixels,
            ),
            max_chars=max_chars,
        )
    parser = PARSERS.get(media_type)
    if parser is None:
        raise DocumentError("unsupported_file_type", "Unsupported file type")
    return _normalize_text(
        parser.parse(content, pdf_max_pages=pdf_max_pages),
        max_chars=max_chars,
    )
