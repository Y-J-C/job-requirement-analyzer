from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.parsing.errors import DocumentError
from app.parsing.parsers import parse_document
from app.parsing.validation import DOCX_MEDIA_TYPE, MARKDOWN_MEDIA_TYPE, PDF_MEDIA_TYPE


def make_pdf(text: str | None = None, *, pages: int = 1, encrypted: bool = False) -> bytes:
    writer = PdfWriter()
    for index in range(pages):
        page = writer.add_blank_page(width=300, height=300)
        if text is not None and index == 0:
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 50 250 Td ({text}) Tj ET".encode("ascii"))
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Contents")] = writer._add_object(stream)
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject(
                        {NameObject("/F1"): writer._add_object(font)}
                    )
                }
            )
    if encrypted:
        writer.encrypt("secret")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("负责数据分析")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "技能"
    table.cell(0, 1).text = "SQL"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_markdown_parser_preserves_markup_as_plain_text() -> None:
    text = parse_document(
        b"# JD\r\n<script>alert(1)</script>\r\nSQL",
        MARKDOWN_MEDIA_TYPE,
        max_chars=100_000,
        pdf_max_pages=50,
    )

    assert text == "# JD\n<script>alert(1)</script>\nSQL"


def test_docx_parser_extracts_paragraphs_and_table_cells() -> None:
    text = parse_document(
        make_docx(), DOCX_MEDIA_TYPE, max_chars=100_000, pdf_max_pages=50
    )

    assert "负责数据分析" in text
    assert "技能\tSQL" in text


def test_pdf_parser_extracts_text_layer() -> None:
    text = parse_document(
        make_pdf("SQL required"), PDF_MEDIA_TYPE, max_chars=100_000, pdf_max_pages=50
    )

    assert text == "SQL required"


@pytest.mark.parametrize(
    ("content", "pdf_max_pages", "code"),
    [
        (make_pdf(encrypted=True), 50, "encrypted_document"),
        (make_pdf(pages=2), 1, "pdf_page_limit_exceeded"),
        (make_pdf(), 50, "no_extractable_text"),
    ],
)
def test_pdf_parser_returns_stable_errors(
    content: bytes, pdf_max_pages: int, code: str
) -> None:
    with pytest.raises(DocumentError) as raised:
        parse_document(content, PDF_MEDIA_TYPE, max_chars=100_000, pdf_max_pages=pdf_max_pages)

    assert raised.value.code == code


def test_parser_rejects_extracted_text_over_limit() -> None:
    with pytest.raises(DocumentError) as raised:
        parse_document(b"123456", MARKDOWN_MEDIA_TYPE, max_chars=5, pdf_max_pages=50)

    assert raised.value.code == "extracted_text_too_large"
